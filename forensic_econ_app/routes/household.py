from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file
from flask_login import login_required, current_user
from ..models.models import db, Evaluee, HouseholdServicesScenario, HouseholdServiceStage
from ..utils.cross_module_data import get_shared_defaults, suggest_scenario_name, get_consistent_dates
from decimal import Decimal
import pandas as pd
import os
from datetime import datetime, date
from openpyxl.utils import get_column_letter
from openpyxl.styles import Border, Side, Alignment, Font

household = Blueprint('household', __name__)

# Helper function to get annual value for a specific year in a scenario
def get_annual_value(scenario, year):
    """Calculate the annual value for a specific year in a scenario."""
    if not scenario or not scenario.stages:
        return 0

    current_year = 1
    for stage in sorted(scenario.stages, key=lambda s: s.stage_number):
        if current_year <= year <= current_year + stage.years - 1:
            # This year is in this stage
            year_in_stage = year - current_year + 1
            growth_factor = (1 + scenario.growth_rate) ** (year_in_stage - 1)
            return float(stage.annual_value * growth_factor)
        current_year += stage.years

    return 0

@household.route('/household/<int:evaluee_id>')
@login_required
def household_form(evaluee_id):
    """Display the household services form and list existing scenarios."""
    evaluee = Evaluee.query.get_or_404(evaluee_id)
    scenarios = HouseholdServicesScenario.query.filter_by(evaluee_id=evaluee_id).all()
    
    # Get smart defaults and cross-module data
    defaults = get_shared_defaults(evaluee_id, 'household')
    consistent_dates = get_consistent_dates(evaluee_id)
    suggested_name = suggest_scenario_name(evaluee_id, 'household', len(scenarios))
    
    return render_template('household/form.html', 
                         evaluee=evaluee, 
                         scenarios=scenarios,
                         defaults=defaults,
                         consistent_dates=consistent_dates,
                         suggested_name=suggested_name,
                         today=date.today().isoformat())

@household.route('/household/<int:evaluee_id>', methods=['POST'])
@login_required
def create_scenario(evaluee_id):
    """Create a new household services scenario."""
    evaluee = Evaluee.query.get_or_404(evaluee_id)

    try:
        # Validate inputs
        scenario_name = request.form.get('scenario_name', '').strip()
        if not scenario_name:
            flash('Scenario name is required.', 'danger')
            return redirect(url_for('household.household_form', evaluee_id=evaluee_id))
        
        # Validate numeric inputs
        try:
            area_wage_adjustment = Decimal(request.form['area_wage_adjustment']) / 100
            reduction_percentage = Decimal(request.form['reduction_percentage']) / 100
            growth_rate = Decimal(request.form['growth_rate']) / 100
            discount_rate = Decimal(request.form['discount_rate']) / 100
        except (ValueError, KeyError) as e:
            flash(f'Invalid numeric input: {str(e)}', 'danger')
            return redirect(url_for('household.household_form', evaluee_id=evaluee_id))
        
        # Validate ranges
        if not (0 <= area_wage_adjustment <= 10):  # 0% to 1000%
            flash('Area wage adjustment must be between 0% and 1000%.', 'danger')
            return redirect(url_for('household.household_form', evaluee_id=evaluee_id))
        
        if not (0 <= reduction_percentage <= 1):  # 0% to 100%
            flash('Reduction percentage must be between 0% and 100%.', 'danger')
            return redirect(url_for('household.household_form', evaluee_id=evaluee_id))
        
        if not (-0.05 <= growth_rate <= 0.20):  # -5% to 20%
            flash('Growth rate must be between -5% and 20%.', 'danger')
            return redirect(url_for('household.household_form', evaluee_id=evaluee_id))
        
        if not (0 <= discount_rate <= 0.20):  # 0% to 20%
            flash('Discount rate must be between 0% and 20%.', 'danger')
            return redirect(url_for('household.household_form', evaluee_id=evaluee_id))

        scenario = HouseholdServicesScenario(
            evaluee_id=evaluee_id,
            scenario_name=scenario_name,
            area_wage_adjustment=area_wage_adjustment,
            reduction_percentage=reduction_percentage,
            growth_rate=growth_rate,
            discount_rate=discount_rate
        )

        db.session.add(scenario)
        db.session.commit()

        flash('Scenario created successfully. Add stages to complete the calculation.')
        return redirect(url_for('household.manage_stages', evaluee_id=evaluee_id, scenario_id=scenario.id))

    except Exception as e:
        db.session.rollback()
        flash(f'Error creating scenario: {str(e)}')
        return redirect(url_for('household.household_form', evaluee_id=evaluee_id))

@household.route('/household/<int:evaluee_id>/scenario/<int:scenario_id>')
@login_required
def view_scenario(evaluee_id, scenario_id):
    """View a specific household services scenario."""
    evaluee = Evaluee.query.get_or_404(evaluee_id)
    scenario = HouseholdServicesScenario.query.get_or_404(scenario_id)

    if scenario.evaluee_id != evaluee_id:
        flash('Access denied.')
        return redirect(url_for('household.household_form', evaluee_id=evaluee_id))

    # Calculate present value
    scenario.present_value = scenario.calculate_present_value()
    db.session.commit()

    return render_template('household/view_scenario.html', evaluee=evaluee, scenario=scenario, Decimal=Decimal)

@household.route('/household/<int:evaluee_id>/scenario/<int:scenario_id>/stages')
@login_required
def manage_stages(evaluee_id, scenario_id):
    """Manage stages for a household services scenario."""
    evaluee = Evaluee.query.get_or_404(evaluee_id)
    scenario = HouseholdServicesScenario.query.get_or_404(scenario_id)

    if scenario.evaluee_id != evaluee_id:
        flash('Access denied.')
        return redirect(url_for('household.household_form', evaluee_id=evaluee_id))

    return render_template('household/stages.html', evaluee=evaluee, scenario=scenario)

@household.route('/household/<int:evaluee_id>/scenario/<int:scenario_id>/stages/add', methods=['POST'])
@login_required
def add_stage(evaluee_id, scenario_id):
    """Add a new stage to a scenario."""
    scenario = HouseholdServicesScenario.query.get_or_404(scenario_id)

    if scenario.evaluee_id != evaluee_id:
        flash('Access denied.')
        return redirect(url_for('household.household_form', evaluee_id=evaluee_id))

    try:
        # Validate inputs
        try:
            stage_number = int(request.form['stage_number'])
            years = int(request.form['years'])
            annual_value = Decimal(request.form['annual_value'])
        except (ValueError, KeyError) as e:
            flash(f'Invalid input: {str(e)}', 'danger')
            return redirect(url_for('household.manage_stages', evaluee_id=evaluee_id, scenario_id=scenario_id))
        
        # Validate ranges
        if not (1 <= stage_number <= 20):
            flash('Stage number must be between 1 and 20.', 'danger')
            return redirect(url_for('household.manage_stages', evaluee_id=evaluee_id, scenario_id=scenario_id))
        
        if not (1 <= years <= 100):
            flash('Years must be between 1 and 100.', 'danger')
            return redirect(url_for('household.manage_stages', evaluee_id=evaluee_id, scenario_id=scenario_id))
        
        if not (0 <= annual_value <= 1000000):
            flash('Annual value must be between $0 and $1,000,000.', 'danger')
            return redirect(url_for('household.manage_stages', evaluee_id=evaluee_id, scenario_id=scenario_id))
        
        # Check for duplicate stage numbers
        existing_stage = HouseholdServiceStage.query.filter_by(
            scenario_id=scenario_id, 
            stage_number=stage_number
        ).first()
        if existing_stage:
            flash(f'Stage number {stage_number} already exists. Please use a different number.', 'danger')
            return redirect(url_for('household.manage_stages', evaluee_id=evaluee_id, scenario_id=scenario_id))

        stage = HouseholdServiceStage(
            scenario_id=scenario_id,
            stage_number=stage_number,
            years=years,
            annual_value=annual_value
        )

        db.session.add(stage)
        db.session.commit()

        # Recalculate present value
        scenario.present_value = scenario.calculate_present_value()
        db.session.commit()

        flash('Stage added successfully.')
    except Exception as e:
        db.session.rollback()
        flash(f'Error adding stage: {str(e)}')

    return redirect(url_for('household.manage_stages', evaluee_id=evaluee_id, scenario_id=scenario_id))

@household.route('/household/<int:evaluee_id>/scenario/<int:scenario_id>/stages/<int:stage_id>/delete', methods=['POST'])
@login_required
def delete_stage(evaluee_id, scenario_id, stage_id):
    """Delete a stage from a scenario."""
    stage = HouseholdServiceStage.query.get_or_404(stage_id)
    scenario = HouseholdServicesScenario.query.get_or_404(scenario_id)

    if scenario.evaluee_id != evaluee_id:
        flash('Access denied.')
        return redirect(url_for('household.household_form', evaluee_id=evaluee_id))

    try:
        db.session.delete(stage)

        # Recalculate present value
        scenario.present_value = scenario.calculate_present_value()
        db.session.commit()

        flash('Stage deleted successfully.')
    except Exception as e:
        db.session.rollback()
        flash(f'Error deleting stage: {str(e)}')

    return redirect(url_for('household.manage_stages', evaluee_id=evaluee_id, scenario_id=scenario_id))

@household.route('/household/<int:evaluee_id>/scenario/<int:scenario_id>/edit', methods=['GET', 'POST'])
@login_required
def edit_scenario(evaluee_id, scenario_id):
    """Edit a household services scenario."""
    evaluee = Evaluee.query.get_or_404(evaluee_id)
    scenario = HouseholdServicesScenario.query.get_or_404(scenario_id)

    if scenario.evaluee_id != evaluee_id:
        flash('Access denied.')
        return redirect(url_for('household.household_form', evaluee_id=evaluee_id))

    if request.method == 'POST':
        try:
            scenario.scenario_name = request.form['scenario_name']
            scenario.area_wage_adjustment = Decimal(request.form['area_wage_adjustment']) / 100
            scenario.reduction_percentage = Decimal(request.form['reduction_percentage']) / 100
            scenario.growth_rate = Decimal(request.form['growth_rate']) / 100
            scenario.discount_rate = Decimal(request.form['discount_rate']) / 100

            # Recalculate present value
            scenario.present_value = scenario.calculate_present_value()
            db.session.commit()

            flash('Scenario updated successfully.')
            return redirect(url_for('household.view_scenario', evaluee_id=evaluee_id, scenario_id=scenario_id))

        except Exception as e:
            db.session.rollback()
            flash(f'Error updating scenario: {str(e)}')

    return render_template('household/edit_scenario.html', evaluee=evaluee, scenario=scenario)

@household.route('/household/<int:evaluee_id>/scenario/<int:scenario_id>/delete', methods=['POST'])
@login_required
def delete_scenario(evaluee_id, scenario_id):
    """Delete a household services scenario."""
    scenario = HouseholdServicesScenario.query.get_or_404(scenario_id)

    if scenario.evaluee_id != evaluee_id:
        flash('Access denied.')
        return redirect(url_for('household.household_form', evaluee_id=evaluee_id))

    try:
        db.session.delete(scenario)
        db.session.commit()
        flash('Scenario deleted successfully.')
    except Exception as e:
        db.session.rollback()
        flash(f'Error deleting scenario: {str(e)}')

    return redirect(url_for('household.household_form', evaluee_id=evaluee_id))

@household.route('/household/<int:evaluee_id>/scenario/<int:scenario_id>/export')
@login_required
def export_scenario(evaluee_id, scenario_id):
    """Export scenario to Excel with corrected annual breakdown and PV calculations."""
    scenario = HouseholdServicesScenario.query.get_or_404(scenario_id)
    evaluee = Evaluee.query.get_or_404(evaluee_id)

    if scenario.evaluee_id != evaluee_id:
        flash('Access denied.')
        return redirect(url_for('household.household_form', evaluee_id=evaluee_id))

    try:
        # Create Excel file
        temp_dir = os.path.dirname(os.path.abspath(__file__))
        temp_path = os.path.join(temp_dir, f'household_services_{scenario_id}.xlsx')

        # Create DataFrame for stages
        stages_df = pd.DataFrame([{
            'Stage': stage.stage_number,
            'Years': stage.years,
            'Annual Value': float(stage.annual_value)
        } for stage in sorted(scenario.stages, key=lambda x: x.stage_number)])

        # -------------------------------
        # Corrected Annual Breakdown
        # -------------------------------
        breakdown_data = []
        total_pv = 0.0
        global_year = 1  # Track the overall year across all stages

        # For each stage, calculate values for each year in that stage
        for stage in sorted(scenario.stages, key=lambda x: x.stage_number):
            for local_year in range(1, stage.years + 1):
                # Growth from the stage's base annual_value
                grown_value = float(stage.annual_value) * (
                    (1 + float(scenario.growth_rate)) ** (local_year - 1)
                )

                # Apply area wage adjustment & reduction
                adjusted_value = grown_value * float(scenario.area_wage_adjustment) * float(scenario.reduction_percentage)

                # Discount from the global year (not local year)
                present_value = adjusted_value / (
                    (1 + float(scenario.discount_rate)) ** global_year
                )

                total_pv += present_value

                breakdown_data.append({
                    'Stage': stage.stage_number,
                    'Stage Year': local_year,
                    'Global Year': global_year,
                    'Base Annual Value': float(stage.annual_value),
                    'Grown Value': grown_value,
                    'Adjusted Value': adjusted_value,
                    'Present Value': present_value
                })
                
                global_year += 1  # Increment global year counter

        # The final present value is simply the sum of discounted, adjusted cash flows
        final_pv = total_pv

        breakdown_df = pd.DataFrame(breakdown_data)

        # Create Excel writer
        with pd.ExcelWriter(temp_path, engine='openpyxl') as writer:
            # Write summary
            summary_df = pd.DataFrame([{
                'Scenario Name': scenario.scenario_name,
                'Area Wage Adjustment': f"{float(scenario.area_wage_adjustment) * 100:.1f}%",
                'Reduction Percentage': f"{float(scenario.reduction_percentage) * 100:.1f}%",
                'Growth Rate': f"{float(scenario.growth_rate) * 100:.1f}%",
                'Discount Rate': f"{float(scenario.discount_rate) * 100:.1f}%",
                'Present Value': final_pv
            }])
            summary_df.to_excel(writer, sheet_name='Summary', index=False)

            # Format the Summary sheet
            worksheet = writer.sheets['Summary']
            for cell in worksheet['F']:  # Present Value column
                if cell.row > 1:  # Skip header
                    cell.number_format = '"$"#,##0.00'

            # Write stages with currency formatting
            stages_df.to_excel(writer, sheet_name='Stages', index=False)
            worksheet = writer.sheets['Stages']
            for cell in worksheet['C']:  # Annual Value column
                if cell.row > 1:  # Skip header
                    cell.number_format = '"$"#,##0.00'

            # Write annual breakdown with currency formatting
            breakdown_df.to_excel(writer, sheet_name='Annual Breakdown', index=False)
            worksheet = writer.sheets['Annual Breakdown']

            # Format currency columns in Annual Breakdown
            currency_columns = ['Base Annual Value', 'Grown Value', 'Adjusted Value', 'Present Value']
            for col_name in currency_columns:
                col_letter = get_column_letter(breakdown_df.columns.get_loc(col_name) + 1)
                for cell in worksheet[col_letter]:
                    if cell.row > 1:  # Skip header
                        cell.number_format = '"$"#,##0.00'

            # Add borders and center alignment to all sheets
            for sheet_name in ['Summary', 'Stages', 'Annual Breakdown']:
                worksheet = writer.sheets[sheet_name]
                for row in worksheet.iter_rows():
                    for cell in row:
                        cell.border = Border(
                            left=Side(style='thin'),
                            right=Side(style='thin'),
                            top=Side(style='thin'),
                            bottom=Side(style='thin')
                        )
                        cell.alignment = Alignment(horizontal='center')

                # Make headers bold
                for cell in worksheet[1]:
                    cell.font = Font(bold=True)

        return send_file(
            temp_path,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'household_services_{evaluee.last_name}_{evaluee.first_name}.xlsx'
        )

    except Exception as e:
        flash(f'Error exporting scenario: {str(e)}')
        return redirect(url_for('household.view_scenario', evaluee_id=evaluee_id, scenario_id=scenario_id))

    finally:
        # Clean up temporary file
        try:
            if os.path.exists(temp_path):
                os.remove(temp_path)
        except Exception:
            pass

@household.route('/household/<int:evaluee_id>/scenario/<int:scenario_id>/print')
@login_required
def print_report(evaluee_id, scenario_id):
    """Generate a printable report for a household services scenario."""
    evaluee = Evaluee.query.get_or_404(evaluee_id)
    scenario = HouseholdServicesScenario.query.get_or_404(scenario_id)

    if scenario.evaluee_id != evaluee_id:
        flash('Access denied.')
        return redirect(url_for('household.household_form', evaluee_id=evaluee_id))

    # Calculate present value
    scenario.present_value = scenario.calculate_present_value()
    db.session.commit()

    return render_template(
        'household/print_report.html',
        evaluee=evaluee,
        scenario=scenario,
        now=datetime.now(),
        Decimal=Decimal
    )

@household.route('/household/<int:evaluee_id>/compare')
@login_required
def compare_scenarios(evaluee_id):
    """Compare two household services scenarios."""
    evaluee = Evaluee.query.get_or_404(evaluee_id)
    all_scenarios = HouseholdServicesScenario.query.filter_by(evaluee_id=evaluee_id).all()

    # Get scenario IDs from query parameters
    scenario1_id = request.args.get('scenario1', type=int)
    scenario2_id = request.args.get('scenario2', type=int)

    scenario1 = None
    scenario2 = None

    if scenario1_id and scenario2_id:
        scenario1 = HouseholdServicesScenario.query.get_or_404(scenario1_id)
        scenario2 = HouseholdServicesScenario.query.get_or_404(scenario2_id)

        # Verify scenarios belong to this evaluee
        if scenario1.evaluee_id != evaluee_id or scenario2.evaluee_id != evaluee_id:
            flash('Access denied.')
            return redirect(url_for('household.household_form', evaluee_id=evaluee_id))

        # Calculate present values
        scenario1.present_value = scenario1.calculate_present_value()
        scenario2.present_value = scenario2.calculate_present_value()
        db.session.commit()

    return render_template(
        'household/compare_scenarios.html',
        evaluee=evaluee,
        all_scenarios=all_scenarios,
        scenario1=scenario1,
        scenario2=scenario2,
        get_annual_value=get_annual_value,
        now=datetime.now()
    )